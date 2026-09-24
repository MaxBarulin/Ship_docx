# -*- coding: utf-8 -*-
"""Пояснительная записка в Word и PDF по шаблону команды.

    python scripts/пз_docx.py            # → docs/проект/ПЗ/Без границ_ПЗ.docx и .pdf

Шаблон - docs/команда/правки 23.09.26/ПЗ_УЖЦ_2026.docx: титульный лист берётся как есть, оформление - как в
шаблоне (А4, поля 30/15/20/20, двойная рамка страницы, Times New Roman, полуторный интервал, логотип команды
внизу по центру, номер страницы справа). Текст - docs/проект/записка.md без шапки для вёрстки. Каждый блок -
свой раздел с новой страницы, внизу слева номер и наименование блока (приложение к КЗ, п. 8). Содержание - поле
оглавления, его, номера страниц и PDF делает Word (COM). Без Word - только .docx, оглавление обновится при открытии.
"""
import copy, io, os, re, subprocess, sys, tempfile
ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
import docx
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor
from PIL import Image

ШАБЛОН = os.path.join(ROOT, "docs", "команда", "правки 23.09.26", "ПЗ_УЖЦ_2026.docx")
ЗАПИСКА = os.path.join(ROOT, "docs", "проект", "записка.md")
КОМАНДА = "Без границ"
OUT_DIR = os.path.join(ROOT, "docs", "проект", "ПЗ")
DOCX = os.path.join(OUT_DIR, "%s_ПЗ.docx" % КОМАНДА)
PDF = os.path.join(OUT_DIR, "%s_ПЗ.pdf" % КОМАНДА)
ШРИФТ = "Times New Roman"
КЕГЛЬ = 12
ИНТЕРВАЛ = 1.5
ШИРИНА_МАКС = 165.0            # мм - ширина поля набора
ВЫСОТА_МАКС = 205.0            # мм - рисунок с подписью помещается на страницу
ПИКС_МАКС = 2000               # длинная сторона растра в документе - около 300 dpi на 165 мм


# --- стили -----------------------------------------------------------------------------------------
def _шрифт(style_or_run, size=None, bold=None):
    f = style_or_run.font
    f.name = ШРИФТ
    rpr = style_or_run.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), ШРИФТ)
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    f.color.rgb = RGBColor(0, 0, 0)


def _стиль(d, имя, основа="Normal"):
    try:
        s = d.styles[имя]
    except KeyError:
        s = d.styles.add_style(имя, WD_STYLE_TYPE.PARAGRAPH)
    s.base_style = d.styles[основа]
    return s


def стили(d):
    т = _стиль(d, "ПЗ Текст")
    _шрифт(т, КЕГЛЬ, False)
    pf = т.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Cm(1.25)
    pf.line_spacing = ИНТЕРВАЛ
    pf.space_before = pf.space_after = Pt(0)
    pf.widow_control = True
    for имя, размер, уровень, до in (("Heading 1", 14, 0, 0), ("Heading 2", 13, 1, 12), ("Heading 3", 12, 2, 8)):
        h = _стиль(d, имя, "ПЗ Текст")
        _шрифт(h, размер, True)
        h.font.italic = False
        pf = h.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.first_line_indent = Cm(1.25)
        pf.space_before = Pt(до)
        pf.space_after = Pt(6)
        pf.keep_with_next = True
        ppr = h.element.get_or_add_pPr()
        for old in ppr.findall(qn("w:outlineLvl")):
            ppr.remove(old)
        ol = OxmlElement("w:outlineLvl")
        ol.set(qn("w:val"), str(уровень))
        ppr.append(ol)
    for имя, выр in (("ПЗ Подпись таблицы", WD_ALIGN_PARAGRAPH.LEFT), ("ПЗ Подпись рисунка", WD_ALIGN_PARAGRAPH.CENTER),
                     ("ПЗ Рисунок", WD_ALIGN_PARAGRAPH.CENTER)):
        s = _стиль(d, имя, "ПЗ Текст")
        s.paragraph_format.alignment = выр
        s.paragraph_format.first_line_indent = Cm(0)
        s.paragraph_format.line_spacing = 1.0
    d.styles["ПЗ Подпись таблицы"].paragraph_format.space_before = Pt(6)
    d.styles["ПЗ Подпись таблицы"].paragraph_format.keep_with_next = True
    d.styles["ПЗ Подпись рисунка"].paragraph_format.space_after = Pt(8)
    d.styles["ПЗ Рисунок"].paragraph_format.keep_with_next = True
    d.styles["ПЗ Рисунок"].paragraph_format.space_before = Pt(6)
    яч = _стиль(d, "ПЗ Таблица", "ПЗ Текст")
    _шрифт(яч, 11, False)
    яч.paragraph_format.first_line_indent = Cm(0)
    яч.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    яч.paragraph_format.line_spacing = 1.0
    сод = _стиль(d, "ПЗ Содержание", "ПЗ Текст")
    _шрифт(сод, 14, True)
    сод.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    сод.paragraph_format.first_line_indent = Cm(0)
    сод.paragraph_format.space_after = Pt(12)
    for уровень in (1, 2, 3):             # строки оглавления - тем же шрифтом (встроенные стили Word «toc N»)
        s = next((x for x in d.styles if x.name and x.name.lower() == "toc %d" % уровень), None)
        if s is None:
            s = d.styles.add_style("toc %d" % уровень, WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = d.styles["ПЗ Текст"]
        _шрифт(s, КЕГЛЬ, уровень == 1)
        s.paragraph_format.first_line_indent = Cm(0)
        s.paragraph_format.left_indent = Cm(0.6 * (уровень - 1))
        s.paragraph_format.line_spacing = 1.15
        s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


# --- поля Word -------------------------------------------------------------------------------------
def _поле(par, код, текст="1", size=None):
    def r(el):
        run = par.add_run()
        _шрифт(run, size or КЕГЛЬ)
        run._r.append(el)
        return run
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); r(b)
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = " %s " % код; r(it)
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate"); r(s)
    run = par.add_run(текст); _шрифт(run, size or КЕГЛЬ)
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r(e)


# --- колонтитул ------------------------------------------------------------------------------------
def _без_рамок(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for края in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:%s" % края); e.set(qn("w:val"), "nil"); b.append(e)
    tblPr.append(b)


def колонтитул(footer, слева, логотип):
    """Внизу страницы: слева номер и наименование блока, по центру логотип команды, справа номер страницы."""
    el = footer._element
    for ch in list(el):
        el.remove(ch)
    tbl = footer.add_table(rows=1, cols=3, width=Mm(ШИРИНА_МАКС))
    _без_рамок(tbl)
    for c, w in zip(tbl.rows[0].cells, (Mm(78), Mm(40), Mm(47))):
        c.width = w
    c0, c1, c2 = tbl.rows[0].cells
    p = c0.paragraphs[0]; p.style = "ПЗ Таблица"
    run = p.add_run(слева); _шрифт(run, 11)
    p = c1.paragraphs[0]; p.style = "ПЗ Таблица"; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(логотип, height=Mm(14))
    p = c2.paragraphs[0]; p.style = "ПЗ Таблица"; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _поле(p, "PAGE", "1", 12)
    for c in (c0, c1, c2):
        tcPr = c._tc.get_or_add_tcPr()
        va = OxmlElement("w:vAlign"); va.set(qn("w:val"), "bottom"); tcPr.append(va)
    footer.add_paragraph().paragraph_format.space_after = Pt(0)


# --- разметка --------------------------------------------------------------------------------------
_ЖИР = re.compile(r"(\*\*[^*]+\*\*)")


def _текст(par, s, size=None, bold=False):
    for кусок in _ЖИР.split(s):
        if not кусок:
            continue
        ж = кусок.startswith("**") and кусок.endswith("**")
        run = par.add_run(кусок[2:-2] if ж else кусок)
        if size or bold or ж:
            _шрифт(run, size, bold or ж)


def _растр(путь, tmp):
    """Растр для документа - длинная сторона не больше ПИКС_МАКС, фото - JPEG, схемы и чертежи - PNG."""
    im = Image.open(путь)
    w, h = im.size
    k = min(1.0, ПИКС_МАКС / float(max(w, h)))
    имя = os.path.join(tmp, "%04d%s" % (len(os.listdir(tmp)), os.path.splitext(путь)[1].lower()))
    if k < 1.0:
        im = im.resize((int(w * k), int(h * k)), Image.LANCZOS)
    # фото (рендеры) - JPEG, схемы и чертежи с малым числом цветов - PNG: линии и текст не мылятся
    фото = имя.endswith(".png") and len(im.convert("RGB").resize((200, 200)).getcolors(40000) or range(40001)) > 6000
    if имя.endswith(".png") and not фото:
        im.save(имя, optimize=True)
    else:
        имя = os.path.splitext(имя)[0] + ".jpg"
        im.convert("RGB").save(имя, quality=90)
    return имя, im.size


def _рисунок(d, путь, tmp):
    файл, (w, h) = _растр(путь, tmp)
    ширина = ШИРИНА_МАКС
    if ширина * h / float(w) > ВЫСОТА_МАКС:
        ширина = ВЫСОТА_МАКС * w / float(h)
    p = d.add_paragraph(style="ПЗ Рисунок")
    p.add_run().add_picture(файл, width=Mm(ширина))


ЗНАК_ММ = 2.15                 # средняя ширина знака Times New Roman 11 pt, мм
ЗНАК_ЖИР_ММ = 2.45             # то же жирным - шапка таблицы
ПОЛЯ_ЯЧЕЙКИ = 2.6              # поля ячейки 1 мм слева и справа и запас, мм
ПОЛЕ_ЯЧЕЙКИ_DXA = 57           # 1 мм


def _ширины(ряды, n):
    """Ширины столбцов, мм: не уже самого длинного слова столбца (шапка - жирным), остаток поля набора - по длине
    текста."""
    мин, хочу = [], []
    for j in range(n):
        тексты = [(r[j] if j < len(r) else "").replace("**", "") for r in ряды]
        шапка = max((len(w) for w in тексты[0].split()), default=1) * ЗНАК_ЖИР_ММ
        тело = max((len(w) for t in тексты[1:] for w in t.split()), default=1) * ЗНАК_ММ
        строка = min(max((len(t) for t in тексты), default=1), 42)
        мин.append(max(шапка, тело) + ПОЛЯ_ЯЧЕЙКИ)
        хочу.append(max(мин[-1], строка * ЗНАК_ММ + ПОЛЯ_ЯЧЕЙКИ))
    if sum(мин) >= ШИРИНА_МАКС:           # длинные слова не рвутся - таблица шире поля набора до 175 мм
        k = min(1.0, 175.0 / sum(мин))
        return [w * k for w in мин]
    if sum(хочу) <= ШИРИНА_МАКС:
        return [w * ШИРИНА_МАКС / sum(хочу) for w in хочу]
    запас = ШИРИНА_МАКС - sum(мин)
    лишнее = sum(х - м for х, м in zip(хочу, мин))
    return [м + запас * (х - м) / лишнее for х, м in zip(хочу, мин)]


def _таблица(d, строки):
    ряды = [[c.strip() for c in r.strip().strip("|").split("|")] for r in строки if not re.match(r"^\|\s*-", r)]
    n = max(len(r) for r in ряды)
    tbl = d.add_table(rows=len(ряды), cols=n)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = d.styles["Table Grid"] if "Table Grid" in [s.name for s in d.styles] else None
    if tbl.style is None:
        tblPr = tbl._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for края in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement("w:%s" % края)
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "000000")
            b.append(e)
        tblPr.append(b)
    ширины = _ширины(ряды, n)
    tblPr = tbl._tbl.tblPr
    tw = tblPr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW"); tblPr.append(tw)
    tw.set(qn("w:type"), "dxa"); tw.set(qn("w:w"), str(int(sum(ширины) * 56.7)))
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    mar = OxmlElement("w:tblCellMar")
    for край in ("left", "right"):
        e = OxmlElement("w:%s" % край); e.set(qn("w:w"), str(ПОЛЕ_ЯЧЕЙКИ_DXA)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tblPr.append(mar)
    for gc, w in zip(tbl._tbl.tblGrid.findall(qn("w:gridCol")), ширины):
        gc.set(qn("w:w"), str(int(w * 56.7)))
    for row in tbl.rows:
        for c, w in zip(row.cells, ширины):
            c.width = Mm(w)
    for i, r in enumerate(ряды):
        row = tbl.rows[i]
        if i == 0:
            trPr = row._tr.get_or_add_trPr()
            th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
        cant = OxmlElement("w:cantSplit"); cant.set(qn("w:val"), "true"); row._tr.get_or_add_trPr().append(cant)
        for j in range(n):
            c = row.cells[j]
            p = c.paragraphs[0]
            p.style = "ПЗ Таблица"
            _текст(p, r[j] if j < len(r) else "", bold=(i == 0))
    return tbl


def _раздел(d, левый_текст, логотип, первый=False):
    """Новый раздел с новой страницы: свой колонтитул, рамка на всех страницах, нумерация сквозная."""
    s = d.add_section(WD_SECTION.NEW_PAGE)
    s.different_first_page_header_footer = False
    sp = s._sectPr
    for pn in sp.findall(qn("w:pgNumType")):
        sp.remove(pn)
    pb = sp.find(qn("w:pgBorders"))
    if pb is not None and pb.get(qn("w:display")):
        del pb.attrib[qn("w:display")]
    s.footer.is_linked_to_previous = False
    колонтитул(s.footer, левый_текст, логотип)
    return s


def разобрать(текст):
    """Строки записки после шапки: блоки начинаются с «## N» и «## Приложение»."""
    строки = текст.split("\n")
    i0 = next(i for i, s in enumerate(строки) if re.match(r"^## 1 ", s))
    return строки[i0:]


def собрать():
    os.makedirs(OUT_DIR, exist_ok=True)
    d = docx.Document(ШАБЛОН)
    body = d.element.body
    дети = list(body)
    for el in дети[2:-1]:                 # титульный лист - первые два абзаца шаблона
        body.remove(el)
    # неиспользуемые картинки шаблона - убрать связи, иначе черновые рисунки уедут в файл
    xml = d.element.xml
    for rid, rel in list(d.part.rels.items()):
        if "image" in rel.reltype and ('"%s"' % rid) not in xml:
            d.part.drop_rel(rid)
    стили(d)
    tmp = tempfile.mkdtemp(prefix="pz_")
    логотип = os.path.join(tmp, "logo.jpeg")
    import zipfile
    with zipfile.ZipFile(ШАБЛОН) as z, io.open(логотип, "wb") as f:
        f.write(z.read("word/media/image7.jpeg"))          # логотип команды из шаблона
    # страница 2 - содержание
    колонтитул(d.sections[0].footer, "Содержание", логотип)
    p = d.add_paragraph("Содержание", style="ПЗ Содержание")
    p = d.add_paragraph(style="ПЗ Текст")
    p.paragraph_format.first_line_indent = Cm(0)
    _поле(p, 'TOC \\o "1-3" \\h \\z \\u', "Оглавление обновляется при открытии")
    строки = разобрать(io.open(ЗАПИСКА, encoding="utf-8").read())
    i = 0
    таблица_буфер = []
    while i < len(строки):
        s = строки[i].rstrip()
        if s.startswith("|"):
            таблица_буфер.append(s)
            i += 1
            if i >= len(строки) or not строки[i].startswith("|"):
                _таблица(d, таблица_буфер)
                таблица_буфер = []
                d.add_paragraph(style="ПЗ Подпись рисунка").paragraph_format.space_after = Pt(4)
            continue
        i += 1
        if not s.strip() or s.startswith("---") or s.startswith(">") or s.startswith("<!--"):
            continue
        m = re.match(r"^(#{2,4}) (.+)$", s)
        if m:
            ур, заг = len(m.group(1)), m.group(2).strip()
            if ур == 2:
                прил = re.match(r"^Приложение ([А-Я])\. (.+)$", заг)
                _раздел(d, (заг if прил else re.sub(r"^(\d+) ", r"\1. ", заг)), логотип)
                if прил:
                    h = d.add_paragraph(style="Heading 1")
                    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    h.paragraph_format.first_line_indent = Cm(0)
                    h.add_run("Приложение %s" % прил.group(1)).add_break(WD_BREAK.LINE)
                    h.add_run(прил.group(2))
                else:
                    d.add_paragraph(заг, style="Heading 1")
            elif ур == 3:
                d.add_paragraph(заг, style="Heading 2")
            else:
                if re.match(r"^[\dА-Я]+\.\d", заг):
                    d.add_paragraph(заг, style="Heading 3")
                else:
                    p = d.add_paragraph(style="ПЗ Текст")
                    p.paragraph_format.keep_with_next = True
                    _текст(p, заг, bold=True)
            continue
        m = re.match(r"^!\[[^\]]*\]\(([^)]+)\)", s)
        if m:
            путь = os.path.normpath(os.path.join(os.path.dirname(ЗАПИСКА), m.group(1)))
            _рисунок(d, путь, tmp)
            continue
        if re.match(r"^Рисунок [\dА-Я]", s):
            d.add_paragraph(s, style="ПЗ Подпись рисунка")
            continue
        if re.match(r"^Таблица [\dА-Я]", s):
            d.add_paragraph(s, style="ПЗ Подпись таблицы")
            continue
        p = d.add_paragraph(style="ПЗ Текст")
        _текст(p, s)
    # последний раздел - приложение: его свойства в конце тела. Первый раздел - титул и содержание
    d.core_properties.title = "Пояснительная записка. Круизное судно «Волжский Горизонт»"
    d.core_properties.author = КОМАНДА
    tmpdoc = DOCX + ".tmp.docx"
    d.save(tmpdoc)
    os.replace(tmpdoc, DOCX)
    return DOCX


def word_pdf():
    """Word: обновить оглавление и номера, сохранить docx, выгрузить PDF, вернуть число страниц."""
    ps = ("$ErrorActionPreference='Stop'; $w = New-Object -ComObject Word.Application; $w.Visible = $false; "
          "$w.DisplayAlerts = 0; $d = $w.Documents.Open('%s'); "
          "foreach ($k in -20, -21, -22) { $st = $d.Styles.Item($k); $st.Font.Name = 'Times New Roman'; $st.Font.Size = 12 }; "
          "$d.Styles.Item(-20).Font.Bold = $true; $d.Fields.Update() | Out-Null; "
          "foreach ($t in $d.TablesOfContents) { $t.Update() }; $d.Repaginate(); "
          "foreach ($t in $d.TablesOfContents) { $t.Update() }; $d.Save(); "
          "$d.ExportAsFixedFormat('%s', 17, $false, 0, 0, 1, 1, 0, $true, $true, 1); "
          "'страниц ' + $d.ComputeStatistics(2); $d.Close(0); $w.Quit()") % (DOCX.replace("'", "''"), PDF.replace("'", "''"))
    r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", ps], capture_output=True, text=True,
                       timeout=900, encoding="cp866", errors="replace")
    return (r.stdout + r.stderr).strip()


def приложения():
    """Приложения отдельными файлами (приложение к КЗ) - «Приложение Б» в папке своего раздела."""
    import shutil
    папка = os.path.join(OUT_DIR, "6 Экономика и финансы")
    os.makedirs(папка, exist_ok=True)
    dst = os.path.join(папка, "Приложение Б.xlsx")
    shutil.copyfile(os.path.join(ROOT, "docs", "проект", "приложение_Б_экономика.xlsx"), dst)
    return dst


if __name__ == "__main__":
    print(os.path.relpath(собрать(), ROOT))
    print(os.path.relpath(приложения(), ROOT))
    if "--без-pdf" not in sys.argv:
        print(word_pdf())
        if os.path.exists(PDF):
            print(os.path.relpath(PDF, ROOT), os.path.getsize(PDF) // 1024, "КБ")
